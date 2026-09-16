"""Genera el PPT de la Reunión Semanal de Marketing a partir de una plantilla
fija (ppt_templates/plantilla_marketing.pptx) y los números ya calculados por
el frontend (mismas funciones que usa el chat de IA del Funnel, _flCalcPeriodo
en templates/index.html) — este módulo solo arma el archivo, no recalcula nada
del negocio.

La plantilla tiene 2 slides:
  - slide 0: portada, con 3 "pills" de ejemplo (proyecto) que se reemplazan por
    N pills según cuántos proyectos entren ese período.
  - slide 1: "master" de un proyecto — se duplica una vez por proyecto y se
    borra al final (nunca queda en el archivo final).
Los índices de shapes de cada slide están documentados abajo porque la
plantilla no usa placeholders nombrados (son Shape/TextBox genéricos de
PowerPoint) — se identifican por su posición en el orden en que fueron creados
en el .pptx original, que es estable mientras no se edite la plantilla.
"""
import copy
import io
import os

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "ppt_templates", "plantilla_marketing.pptx")

# slide 1 (master de proyecto) — índices de shapes en el orden del .pptx
_IDX_TITULO_PROYECTO = 1
_IDX_SUBTITULO_PERIODO = 2
_IDX_BADGE_BG = 3
_IDX_BADGE_TXT = 4
_IDX_KPI1_VAL = 6
_IDX_KPI1_LBL = 7
_IDX_KPI2_VAL = 9
_IDX_KPI2_LBL = 10
_IDX_KPI3_VAL = 12
_IDX_KPI3_LBL = 13
_IDX_TABLA = 15
_IDX_ANALISIS = 19
_IDX_FOOTER_PROYECTO = 23

# slide 0 (portada)
_IDX_PORTADA_PERIODO = 3
_IDX_PORTADA_PILL_BG_1 = 4   # 1er par (fondo, texto) de pill de ejemplo — plantilla de estilo
_IDX_PORTADA_PILL_TXT_1 = 5


def _set_text(shape, texto):
    """Sobreescribe el texto de un shape preservando el formato del primer run
    (fuente/tamaño/color de la plantilla) en vez de resetearlo con .text_frame.text=."""
    tf = shape.text_frame
    if tf.paragraphs and tf.paragraphs[0].runs:
        p0 = tf.paragraphs[0]
        p0.runs[0].text = texto
        for r in p0.runs[1:]:
            r.text = ""
        for p in tf.paragraphs[1:]:
            for r in p.runs:
                r.text = ""
    else:
        tf.text = texto


def _duplicar_slide(prs, slide_origen):
    layout = slide_origen.slide_layout
    nueva = prs.slides.add_slide(layout)
    for shp in list(nueva.shapes):  # limpiar placeholders que trae el layout
        shp._element.getparent().remove(shp._element)
    for shp in slide_origen.shapes:
        nueva.shapes._spTree.append(copy.deepcopy(shp._element))
    return nueva


def _quitar_shape(shapes, idx):
    shp = shapes[idx]
    shp._element.getparent().remove(shp._element)


def _tabla_ajustar_filas(tabla, n_filas_datos):
    """La plantilla trae 3 filas de datos + 1 de encabezado. Clona/borra filas
    para que la tabla tenga exactamente n_filas_datos, sin tocar el encabezado."""
    tbl = tabla._tbl
    trs = tbl.findall(qn("a:tr"))
    n_actual = len(trs) - 1
    if n_filas_datos > n_actual:
        ultima = trs[-1]
        for _ in range(n_filas_datos - n_actual):
            tbl.append(copy.deepcopy(ultima))
    elif n_filas_datos < n_actual:
        trs = tbl.findall(qn("a:tr"))
        for tr in trs[1 + n_filas_datos:]:
            tbl.remove(tr)


def _fmt_prom(n):
    try:
        return f"{float(n):.1f}"
    except (TypeError, ValueError):
        return "—"


def _llenar_slide_proyecto(slide, proyecto, tiene_comparacion):
    shapes = list(slide.shapes)
    nombre = proyecto.get("nombre", "")
    leads_a = proyecto.get("leadsA", 0)
    leads_b = proyecto.get("leadsB")
    prom_a = proyecto.get("promA", 0)
    prom_b = proyecto.get("promB")
    visitas_a = proyecto.get("visitasA", 0)
    pct = proyecto.get("pctCambio")
    canales = proyecto.get("canales", [])[:8]  # tope razonable para que la tabla no se desborde
    analisis = proyecto.get("analisis") or ""

    _set_text(shapes[_IDX_TITULO_PROYECTO], nombre)
    _set_text(shapes[_IDX_SUBTITULO_PERIODO], proyecto.get("periodoLabel", ""))
    _set_text(shapes[_IDX_FOOTER_PROYECTO], nombre.upper())

    if tiene_comparacion and pct is not None:
        flecha = "▲" if pct >= 0 else "▼"
        _set_text(shapes[_IDX_BADGE_TXT], f"{flecha} {abs(pct):.1f}% leads")
    else:
        _quitar_shape(shapes, _IDX_BADGE_TXT)
        _quitar_shape(shapes, _IDX_BADGE_BG)

    if tiene_comparacion and leads_b is not None:
        _set_text(shapes[_IDX_KPI1_VAL], f"{leads_a} vs {leads_b}")
        _set_text(shapes[_IDX_KPI1_LBL], "Leads totales (semana vs. período de comparación)")
        _set_text(shapes[_IDX_KPI2_VAL], f"{_fmt_prom(prom_a)} vs {_fmt_prom(prom_b)}")
    else:
        _set_text(shapes[_IDX_KPI1_VAL], str(leads_a))
        _set_text(shapes[_IDX_KPI1_LBL], "Leads totales de la semana")
        _set_text(shapes[_IDX_KPI2_VAL], _fmt_prom(prom_a))
    _set_text(shapes[_IDX_KPI2_LBL], "Promedio de leads por día")
    _set_text(shapes[_IDX_KPI3_VAL], str(visitas_a))
    _set_text(shapes[_IDX_KPI3_LBL], "Visitas totales generadas en la semana")

    tabla_shape = next(s for s in shapes if s.has_table)
    tabla = tabla_shape.table
    _tabla_ajustar_filas(tabla, max(len(canales), 1))
    if canales:
        for i, c in enumerate(canales):
            fila = tabla.rows[i + 1]
            _set_text_cell(fila.cells[0], str(c.get("canal", "—")))
            _set_text_cell(fila.cells[1], str(c.get("leads", 0)))
            _set_text_cell(fila.cells[2], str(c.get("visitas", 0)))
    else:
        _set_text_cell(tabla.rows[1].cells[0], "Sin canales registrados")
        _set_text_cell(tabla.rows[1].cells[1], "—")
        _set_text_cell(tabla.rows[1].cells[2], "—")

    _set_text(shapes[_IDX_ANALISIS], analisis or "Sin datos suficientes para un análisis automático esta semana.")


def _set_text_cell(cell, texto):
    tf = cell.text_frame
    if tf.paragraphs and tf.paragraphs[0].runs:
        tf.paragraphs[0].runs[0].text = texto
        for r in tf.paragraphs[0].runs[1:]:
            r.text = ""
    else:
        tf.text = texto


def _llenar_portada(slide, periodo_label, nombres_proyectos, slide_width):
    shapes = list(slide.shapes)
    _set_text(shapes[_IDX_PORTADA_PERIODO], periodo_label)

    # Capturamos el estilo del 1er par (fondo + texto) antes de borrar los 3
    # pills de ejemplo, para clonarlo tantas veces como proyectos haya.
    bg_tmpl = copy.deepcopy(shapes[_IDX_PORTADA_PILL_BG_1]._element)
    txt_tmpl = copy.deepcopy(shapes[_IDX_PORTADA_PILL_TXT_1]._element)
    ancho_pill = shapes[_IDX_PORTADA_PILL_BG_1].width
    top_pill = shapes[_IDX_PORTADA_PILL_BG_1].top
    alto_pill = shapes[_IDX_PORTADA_PILL_BG_1].height
    gap = 274320

    # Borra los 3 pills de ejemplo (6 shapes: 3 pares bg+texto), de atrás hacia adelante
    for idx in range(9, 3, -1):
        _quitar_shape(shapes, idx)

    n = max(len(nombres_proyectos), 1)
    total_ancho = n * ancho_pill + (n - 1) * gap
    start_x = (slide_width - total_ancho) // 2

    spTree = slide.shapes._spTree
    for i, nombre in enumerate(nombres_proyectos):
        left = start_x + i * (ancho_pill + gap)
        nuevo_bg = copy.deepcopy(bg_tmpl)
        nuevo_txt = copy.deepcopy(txt_tmpl)
        spTree.append(nuevo_bg)
        spTree.append(nuevo_txt)
        bg_shape = slide.shapes[-2]
        txt_shape = slide.shapes[-1]
        bg_shape.left = Emu(left)
        bg_shape.top = Emu(top_pill)
        txt_shape.left = Emu(left)
        txt_shape.top = Emu(top_pill)
        _set_text(txt_shape, nombre.upper())


def generar_ppt(periodo_label, tiene_comparacion, proyectos):
    """proyectos: lista de dicts con nombre/leadsA/leadsB/promA/promB/visitasA/
    pctCambio/canales/analisis/periodoLabel (ver _llenar_slide_proyecto)."""
    prs = Presentation(TEMPLATE_PATH)
    portada = prs.slides[0]
    master = prs.slides[1]

    _llenar_portada(portada, periodo_label, [p.get("nombre", "") for p in proyectos], prs.slide_width)

    for p in proyectos:
        nueva = _duplicar_slide(prs, master)
        _llenar_slide_proyecto(nueva, p, tiene_comparacion)

    # Borrar el slide master original (índice 1): ya se duplicó una vez por proyecto.
    # drop_rel además de la relationship saca la part física — si solo se saca de
    # _sldIdLst queda huérfana y rompe la numeración de partes al guardar (colisión
    # de nombres ppt/slides/slideN.xml).
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    sld_id_elem = slides[1]
    prs.part.drop_rel(sld_id_elem.rId)
    xml_slides.remove(sld_id_elem)

    out = io.BytesIO()
    prs.save(out)
    out.seek(0)
    return out
